from pathlib import Path

from openpyxl import load_workbook

from app.loaders.base import BaseDocumentLoader
from app.models import DocumentChunk


class ExcelLoader(BaseDocumentLoader):
    """
    Loader for Excel files.
    """

    def load(
        self,
        file_path: str,
    ) -> list[DocumentChunk]:

        path = Path(file_path)

        workbook = load_workbook(
            file_path,
            data_only=True,
        )

        documents = []

        chunk_index = 0

        for worksheet in workbook.worksheets:

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if not values:
                    continue

                content = " | ".join(values)

                chunk_index += 1

                chunk = DocumentChunk(
                    document_id=path.stem,
                    file_name=path.name,
                    content=content,
                    page=None,
                    section=worksheet.title,
                    chunk_index=chunk_index,
                    file_type="xlsx",
                )

                documents.append(chunk)

        return documents